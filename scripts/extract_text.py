#!/usr/bin/env python3
"""Print the text of a .docx or .pdf, for reading a person's existing documents.

  python extract_text.py old-resume.pdf
"""
import sys
from pathlib import Path


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    path = Path(sys.argv[1])
    suffix = path.suffix.lower()
    if suffix == ".docx":
        from docx import Document
        doc = Document(str(path))
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                lines.append("\t".join(cell.text for cell in row.cells))
        print("\n".join(lines))
    elif suffix == ".pdf":
        from pypdf import PdfReader
        print("\n\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages))
    else:
        sys.exit("Give a .docx or .pdf file. Convert .doc or Pages files to .docx first.")


if __name__ == "__main__":
    main()
