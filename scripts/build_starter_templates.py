#!/usr/bin/env python3
"""Rebuild the starter templates in templates/classic/. Maintainers only.

Each template holds one sample paragraph per role, and its style map points at
them, exactly like a template made from a person's own resume.
"""
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

OUT = Path(__file__).resolve().parent.parent / "templates" / "classic"
FONT = "Calibri"


def new_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(0.75))
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(0)
    return doc


def para(doc, text, *, size=10.5, bold=False, align=None, before=0, after=0, style=None):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    run.font.name, run.font.size = FONT, Pt(size)
    if bold:
        run.bold = True
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
    if align:
        p.alignment = align
    return p


def underline_rule(p):
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "444444")):
        bottom.set(qn(f"w:{key}"), value)
    borders.append(bottom)
    p._p.get_or_add_pPr().append(borders)


def save(doc, name, roles, contract):
    OUT.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT / f"{name}.docx"))
    style = {"template": f"{name}.docx", "roles": {r: {"paragraph": i} for i, r in enumerate(roles)}, "contract": contract}
    (OUT / f"{name}.style.json").write_text(json.dumps(style, indent=2) + "\n", encoding="utf-8")


def resume():
    doc = new_document()
    center = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, "Full Name", size=20, bold=True, align=center)
    para(doc, "City, ST | email@example.com | 555-555-5555 | linkedin.com/in/name", size=10, align=center, after=4)
    underline_rule(para(doc, "SECTION HEADING", size=11, bold=True, before=8, after=3))
    para(doc, "A short paragraph, used for the summary.", after=2)
    entry = para(doc, "Employer or Project", bold=True, before=4)
    entry.add_run(" — Role title\tJan 2020 – Present").font.name = FONT
    width = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
    entry.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
    para(doc, "An accomplishment bullet.", style="List Bullet", after=1)
    para(doc, "Label: comma, separated, values", after=1)
    save(doc, "resume", ["name", "contact", "section", "body", "entry", "bullet", "skills"], {"page_limit": 2})


def cover_letter():
    doc = new_document()
    para(doc, "Full Name", size=16, bold=True)
    para(doc, "City, ST | email@example.com | 555-555-5555", size=10, after=14)
    para(doc, "January 1, 2026", after=12)
    para(doc, "Hiring Team, Company Name", after=12)
    para(doc, "Dear Hiring Team,", after=10)
    para(doc, "A body paragraph.", after=10)
    para(doc, "Sincerely,", after=18)
    para(doc, "Full Name")
    save(doc, "cover-letter", ["name", "contact", "date", "recipient", "greeting", "body", "closing", "signature"], {"page_limit": 1})


if __name__ == "__main__":
    resume()
    cover_letter()
    print(f"wrote templates to {OUT}")
